# -*- coding: utf-8 -*-
import json, statistics as st, unicodedata, re, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def norm(n):
    n=unicodedata.normalize('NFD', n); n=''.join(c for c in n if unicodedata.category(c)!='Mn')
    return re.sub(r'\s+',' ',n).strip().lower()

P=json.load(open('projecao.json')); BAR=P['BAR']; alvo=P['alvo']; ref=P['ref']
B={r['nome']:r for r in json.load(open('base_longit.json'))}
Bn={norm(k):v for k,v in B.items()}
VAR={'eduardo lorenco moreira':'eduardo lourenco moreira','maria amelia vozniak deluca schneider':'maria amelia vozniak deluca',
     'amanda de matos da silva':'ananda de matos silva','bruna gonzalez nejm':'bruna gonzales nejm'}
def reg(nome):
    return Bn.get(VAR.get(norm(nome),norm(nome)))

def banda(m):
    if m>=8: return 'Folga'
    if m>=3: return 'Confortável'
    if m>=-3: return 'ZONA DE DECISÃO'
    return 'ABAIXO DA BARRA'

for a in alvo:
    a['banda']=banda(a['margem'])
    r=reg(a['nome'])
    mom=''
    if r and r.get('z26') is not None and r.get('s02_z') is not None:
        d=float(r['s02_z'])-float(r['z26'])
        mom='subiu' if d>0.5 else ('caiu' if d<-0.5 else 'estável')
    a['momento']=mom
    a['turma']='T8 (12ª)' if a['fase']==11 else 'T9 (11ª)'

# PCP por turma e total
def pcp(xs): 
    ac=sum(1 for a in xs if a['margem']>0); return ac,len(xs),100*ac/len(xs)
t8=[a for a in alvo if a['fase']==11]; t9=[a for a in alvo if a['fase']==10]
print('T8:',pcp(t8),' T9:',pcp(t9),' total:',pcp(alvo))

ARIAL='Arial'
HDRF=PatternFill('solid',fgColor='1F4E5F'); HDR=Font(name=ARIAL,size=10,bold=True,color='FFFFFF')
CF=Font(name=ARIAL,size=10); TIT=Font(name=ARIAL,size=14,bold=True,color='1F4E5F')
THIN=Side(style='thin',color='BFBFBF'); BORD=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
ALT=PatternFill('solid',fgColor='F2F7F9'); RED=PatternFill('solid',fgColor='F8D7DA')
AMB=PatternFill('solid',fgColor='FFF3CD'); GRN=PatternFill('solid',fgColor='D4EDDA')
wb=Workbook()

def sheet(name,headers,rows,widths,fills=None,freeze='C2'):
    ws=wb.create_sheet(name); ws.append(headers)
    for c in range(1,len(headers)+1):
        cell=ws.cell(1,c); cell.font=HDR; cell.fill=HDRF
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cell.border=BORD
    for i,row in enumerate(rows,start=2):
        ws.append(row)
        f=fills(row) if fills else None
        for j in range(1,len(headers)+1):
            cell=ws.cell(i,j); cell.font=CF; cell.border=BORD
            if f: cell.fill=f
            elif i%2==0: cell.fill=ALT
    for c,w in enumerate(widths,start=1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes=freeze; ws.row_dimensions[1].height=30
    ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}'
    return ws

ordem={'ABAIXO DA BARRA':0,'ZONA DE DECISÃO':1,'Confortável':2,'Folga':3}
rows=[[a['turma'],a['nome'],round(a['ptotal'],1) if a['ptotal'] is not None else None,
       round(a['rel'],2),round(a['margem'],2),a['banda'],a['momento'],
       'estimado via Simulado 02 (faltou ao TP)' if a.get('estim') else '']
      for a in sorted(alvo,key=lambda a:(ordem[a['banda']],a['margem']))]
sheet('Projecao_individual',
      ['Turma','Nome','TP2026 %','vs rede NAPISUL (p.p.)','Margem sobre a barra (p.p.)','Banda','Momento maio→julho','Obs.'],
      rows,[10,38,10,14,14,17,13,32],
      fills=lambda r: RED if r[5]=='ABAIXO DA BARRA' else (AMB if r[5]=='ZONA DE DECISÃO' else (GRN if r[5]=='Folga' else None)))

rows=[[r2['nome'],r2['fase'],round(r2['rel'],2),'ABAIXO' if r2['rel']<BAR else '']
      for r2 in sorted(ref,key=lambda x:x['rel'])]
sheet('Referencia_T6_T7',['Nome (coorte ENAMED 2025)','Fase no TP2025','vs rede (p.p.)','Posição vs barra'],
      rows,[38,13,12,12],fills=lambda r: RED if r[3]=='ABAIXO' else None,freeze='B2')

ws=wb.create_sheet('Cenarios')
ws.column_dimensions['A'].width=44
for c in 'BCDE': ws.column_dimensions[c].width=13
dados=[
 ('Cenários de PCP (percentual de concluintes proficientes)','','','',''),
 ('','T8 (12ª)','T9 (11ª)','Coorte total','Conceito'),
 ('Base — estado maio/julho', f'{pcp(t8)[2]:.1f}%', f'{pcp(t9)[2]:.1f}%', f'{pcp(alvo)[2]:.1f}%', '4'),
 ('Pessimista — barra 3 p.p. mais alta',
  f'{100*sum(1 for a in t8 if a["margem"]>3)/len(t8):.1f}%',
  f'{100*sum(1 for a in t9 if a["margem"]>3)/len(t9):.1f}%',
  f'{100*sum(1 for a in alvo if a["margem"]>3)/len(alvo):.1f}%','4'),
 ('Otimista — barra 3 p.p. mais baixa',
  f'{100*sum(1 for a in t8 if a["margem"]>-3)/len(t8):.1f}%',
  f'{100*sum(1 for a in t9 if a["margem"]>-3)/len(t9):.1f}%',
  f'{100*sum(1 for a in alvo if a["margem"]>-3)/len(alvo):.1f}%','4–5'),
 ('','','','',''),
 (f'Erro-padrão binomial ≈ ±4,7 p.p. (n={len(alvo)}). Faixas: PCP ≥90% → 5 · 75–90 → 4 · 60–75 → 3.','','','',''),
]
for i,row in enumerate(dados,start=1):
    for j,v in enumerate(row,start=1):
        ws.cell(i,j,v).font=Font(name=ARIAL,size=10,bold=(i in (1,2)))
ws.sheet_view.showGridLines=False

ws=wb.create_sheet('Leia-me',0)
ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=115
ws['A1']='Projeção ENAMED 2026 — T8 (atual 12ª) e T9 (atual 11ª)'; ws['A1'].font=TIT
linhas=[
 ('',''),
 ('O que projeta','O conceito ENAMED do curso vem do PCP: % de concluintes com nota ≥ 60 na escala TRI do INEP. Faixa 4 = PCP entre 75% e 90%. UNIDAVI 2025: PCP 82%, Conceito 4.'),
 ('Âncora','A coorte que fez o ENAMED 2025 (T6+T7, n=53 no TP2025) teve 18% abaixo da barra. Expressando a nota de cada aluno como diferença vs a média da rede NAPISUL da própria fase/ano (o que remove a dificuldade da prova), a barra de proficiência cai em −8,07 p.p. vs rede. Essa barra é aplicada à coorte 2026 (T8+T9) no TP2026.'),
 ('Premissas','1) A rede NAPISUL é um metro estável entre anos (n=528–1.128 por fase). 2) A relação TP→ENAMED da escola é estável entre 2025 e 2026. 3) O crescimento maio→outubro da UNIDAVI acompanha o da rede (o que os simulados sugerem ser conservador: S01→S02 a coorte subiu). 4) Coorte 2026 = 34 (T8) + 31 (T9); 5 alunos sem TP2026 estimados pelo z do Simulado 02; Danton Capistrano Ferreira (T8) sem dado recente — fora da conta.'),
 ('Resultado',f'PCP base {pcp(alvo)[2]:.1f}% (T8 {pcp(t8)[2]:.1f}%, T9 {pcp(t9)[2]:.1f}%) → Conceito 4, com folga sobre o piso de 75%. Conceito 5 (90%) fica a ~5 alunos de distância; o cenário otimista encosta nele.'),
 ('Como usar','A aba Projecao_individual ordena por margem. Os 6 nomes na ZONA DE DECISÃO (±3 p.p. da barra) são onde 1 aluno = 1,5 p.p. de PCP: é o grupo de maior retorno pedagógico até outubro. Os 5 ABAIXO DA BARRA precisam de plano estruturado (todos já eram crônicos no mapa de atenção).'),
 ('Limitações','Não é predição determinística: erro-padrão ±4,7 p.p. só do tamanho da coorte; a composição real de concluintes inscritos pode diferir; mudanças de dificuldade relativa do ENAMED 2026 e a régua do painel de especialistas do INEP são incógnitas. Use como instrumento de priorização, não como nota prevista.'),
]
r=2
for a,b in linhas:
    ws.cell(r,1,a).font=Font(name=ARIAL,size=10,bold=bool(a))
    c=ws.cell(r,2,b); c.font=Font(name=ARIAL,size=10); c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(r,1).alignment=Alignment(vertical='top')
    ws.row_dimensions[r].height=max(15,13*(len(b)//108+1)); r+=1
ws.sheet_view.showGridLines=False
del wb['Sheet']
wb.save('Projecao_ENAMED_2026.xlsx')
print('ok')
