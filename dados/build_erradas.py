# -*- coding: utf-8 -*-
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import imgdata as D
LET='ABCDE'; FASES=['6ª Fase','7ª Fase','10ª Fase','11ª Fase']
N1={'6ª Fase':27,'7ª Fase':33,'10ª Fase':31,'11ª Fase':27}
N2={'6ª Fase':28,'7ª Fase':32,'10ª Fase':31,'11ª Fase':29}
def load(f):
    o={}
    for i,c in enumerate(json.load(open(f)),1):
        o[i]={s['name']:([float(x) if x is not None else 0 for x in s['vals']]+[0.0]*5)[:5] for s in c['series']}
    return o
c1,c2=load('charts1.json'),load('charts2.json')
GAB2=open('gab_s2_oficial.txt').read().strip()
KEY1={q:D.GAB_S1[q-1] for q in range(1,41)}; KEY1[39]='E'
KEY2={q:GAB2[q-1] for q in range(1,61)}
q1=[q for q in range(1,41) if q!=24]

# tema/área curados a partir dos enunciados lidos
TEMA={
('S01',8):('Saúde Coletiva','Segurança do paciente — conceito de evento adverso (RDC 36/2013)'),
('S01',20):('Saúde Coletiva','Pol. Nac. Promoção da Saúde — valores/princípios'),
('S01',17):('Saúde Coletiva','Atenção Primária — função e definição'),
('S01',6):('Saúde Coletiva','Composição da equipe de Saúde Ribeirinha (Port. 837/2014)'),
('S01',11):('Epidemiologia','Delineamento de estudo — transversal x coorte'),
('S01',14):('Saúde Coletiva','Abrangência da RDC 36/2013 (segurança do paciente)'),
('S01',18):('Saúde Coletiva','Princípios do SUS/RAS na Atenção Básica'),
('S01',7):('Saúde Coletiva','Consultório na Rua — registro (SISAB)'),
('S01',37):('Clínica Médica','Arterite de células gigantes — VHS (geriatria)'),
('S01',40):('Pediatria/Infecto','Difteria — clínica das pseudomembranas'),
('S01',5):('Saúde Coletiva','Definição de vigilância em saúde (Res. 588/2018 CNS)'),
('S01',12):('Saúde Coletiva','Lei 8.142/90 — periodicidade da Conferência de Saúde'),
('S02',59):('Saúde Mental','Acolhimento à crise no CAPS'),
('S02',33):('Epidemiologia','Valor preditivo positivo — teorema de Bayes'),
('S02',49):('Gestão/MFC','Organização do processo de trabalho na eSF'),
('S02',48):('Saúde Coletiva/Ética','ECA — recusa vacinal e proteção da criança'),
('S02',57):('Gestão/MFC','Matriciamento — instrumentos do apoio matricial'),
('S02',16):('Saúde Mental','Comportamento compulsivo — abordagem em MFC'),
('S02',46):('Clínica Médica','Derrame pleural/neoplasia — investigação'),
('S02',10):('Gestão/MFC','Modelos de agendamento (acesso avançado)'),
('S02',1):('Saúde Coletiva','Lei 8.142/90 — requisitos p/ repasse (fundo+conselho)'),
('S02',12):('GO/Ética','Gravidez indesejada — conduta em UBS'),
('S02',25):('Saúde Coletiva/Infecto','Vigilância da febre amarela — epizootia'),
('S02',51):('Saúde Coletiva','Política Nacional de Humanização — conceito'),
('S01',27):('Clínica/Geriatria','Idoso frágil com Alzheimer avançado — cuidado'),
('S02',4):('Clínica/Geriatria','Paciente acamado/hemiplégico — cuidado domiciliar'),
('S02',28):('Gestão/MFC','Instrumentos de gestão da eSF'),
('S02',35):('Pediatria','Criança indígena com pneumonia (AIDPI)'),
('S02',53):('Saúde Coletiva','eMulti — composição das equipes multiprofissionais'),
}
def analisa(ch,qs,KEY,N,tag,keyerr):
    tot=sum(N.values()); out=[]
    for i,q in enumerate(qs,1):
        k=LET.index(KEY[q]); p=sum(round(ch[i][f][k]*N[f]) for f in FASES)/tot
        pc=(sum(round(ch[i][f][k]*N[f]) for f in ('10ª Fase','11ª Fase')))/(N['10ª Fase']+N['11ª Fase'])
        w=sorted(((sum(ch[i][f][j]*N[f] for f in FASES)/tot,LET[j]) for j in range(5) if j!=k),reverse=True)
        dp,dl=w[0]
        if q in keyerr: cls='ERRO DE GABARITO'
        elif dp>=0.45 and dp>p: cls='ARMADILHA'
        elif dp>=0.30 and dp>p: cls='CONCEITO TROCADO'
        else: cls='SEM DOMÍNIO'
        area,tm=TEMA.get((tag,q),('—','—'))
        out.append([tag,q,KEY[q],round(p*100,1),round(pc*100,1),dl,round(dp*100,1),cls,area,tm])
    return out
rows=analisa(c1,q1,KEY1,N1,'S01',{17,20})+analisa(c2,list(range(1,61)),KEY2,N2,'S02',{59})

ARIAL='Arial'
HDRF=PatternFill('solid',fgColor='1F4E5F'); HDR=Font(name=ARIAL,size=10,bold=True,color='FFFFFF')
CF=Font(name=ARIAL,size=10); TIT=Font(name=ARIAL,size=14,bold=True,color='1F4E5F')
THIN=Side(style='thin',color='BFBFBF'); BORD=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
ALT=PatternFill('solid',fgColor='F2F7F9'); RED=PatternFill('solid',fgColor='F8D7DA')
AMB=PatternFill('solid',fgColor='FFF3CD'); PUR=PatternFill('solid',fgColor='E7DDF0')
wb=Workbook()
def sheet(name,headers,data,widths,fills=None,freeze='D2'):
    ws=wb.create_sheet(name); ws.append(headers)
    for c in range(1,len(headers)+1):
        cell=ws.cell(1,c); cell.font=HDR; cell.fill=HDRF
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cell.border=BORD
    for i,row in enumerate(data,start=2):
        ws.append(row); f=fills(row) if fills else None
        for j in range(1,len(headers)+1):
            cell=ws.cell(i,j); cell.font=CF; cell.border=BORD
            if f: cell.fill=f
            elif i%2==0: cell.fill=ALT
    for c,w in enumerate(widths,start=1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes=freeze; ws.row_dimensions[1].height=30
    ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(data)+1}'
    return ws
def fill(row):
    if row[7]=='ERRO DE GABARITO': return PUR
    if row[3]<35: return RED
    if row[3]<50: return AMB
    return None
H=['Simulado','Q','Gab.','% acerto\n(geral)','% acerto\ncoorte 10+11','Distrator\ndominante','% no\ndistrator','Tipo de erro','Área ENAMED','Tema específico']
allrows=sorted(rows,key=lambda r:r[3])
sheet('Todas_por_acerto',H,allrows,[9,5,6,9,11,9,8,17,20,52],fills=fill)
crit=[r for r in allrows if r[3]<50]
sheet('Criticas_<50',H,crit,[9,5,6,9,11,9,8,17,20,52],fills=fill)

# resumo por área (só críticas <50)
import collections
byarea=collections.Counter(r[8] for r in crit)
bycls=collections.Counter(r[7] for r in crit)
ws=wb.create_sheet('Resumo')
ws.column_dimensions['A'].width=40; ws.column_dimensions['B'].width=12
r=1
ws.cell(r,1,'Questões com acerto < 50% — por ÁREA').font=Font(name=ARIAL,bold=True,size=12); r+=2
ws.cell(r,1,'Área').font=HDR; ws.cell(r,1).fill=HDRF; ws.cell(r,2,'nº').font=HDR; ws.cell(r,2).fill=HDRF; r+=1
for a,n in byarea.most_common():
    ws.cell(r,1,a).font=CF; ws.cell(r,2,n).font=CF; r+=1
r+=1
ws.cell(r,1,'Por TIPO DE ERRO').font=Font(name=ARIAL,bold=True,size=12); r+=2
ws.cell(r,1,'Tipo').font=HDR; ws.cell(r,1).fill=HDRF; ws.cell(r,2,'nº').font=HDR; ws.cell(r,2).fill=HDRF; r+=1
for a,n in bycls.most_common():
    ws.cell(r,1,a).font=CF; ws.cell(r,2,n).font=CF; r+=1
ws.sheet_view.showGridLines=False

ws=wb.create_sheet('Leia-me',0)
ws.column_dimensions['A'].width=22; ws.column_dimensions['B'].width=116
ws['A1']='Questões mais erradas — Simulados ENAMED 2026 (S01 e S02)'; ws['A1'].font=TIT
linhas=[
 ('',''),
 ('Dado','Distribuição de respostas A–E de cada questão, por fase, extraída do cache dos gráficos dos dashboards. % de acerto = quem marcou o gabarito; % no distrator = alternativa errada mais escolhida.'),
 ('Coorte 10+11','Coluna à parte com o acerto só das fases 10ª e 11ª (a coorte que faz o ENAMED 2026) — é a leitura que importa para priorizar revisão de véspera.'),
 ('Tipo de erro','ERRO DE GABARITO (chave provável equivocada; ver análise anterior) · ARMADILHA (um único distrator leva ≥45%, vence o gabarito — em geral pegadinha ou conceito muito próximo) · CONCEITO TROCADO (um distrator concentra 30–45%) · SEM DOMÍNIO (respostas espalhadas, ninguém sabe).'),
 ('Área/Tema','Classificação manual pelo enunciado, na taxonomia de áreas do ENAMED. Use a aba Resumo para ver a concentração.'),
 ('Achado principal','Das 29 questões com acerto < 50% nos dois simulados, a grande maioria é de Saúde Coletiva / Gestão do SUS / legislação — não de clínica. O gargalo é conteúdo normativo (leis 8.080/8.142, RDCs, portarias, políticas nacionais), não raciocínio clínico.'),
 ('Ressalva','3 das mais erradas são erros de gabarito (S01 Q17, S01 Q20, S02 Q59): o "erro" é da prova, não do aluno. Estão marcadas em roxo e NÃO devem entrar no diagnóstico de conteúdo.'),
]
r=2
for a,b in linhas:
    ws.cell(r,1,a).font=Font(name=ARIAL,size=10,bold=bool(a))
    c=ws.cell(r,2,b); c.font=Font(name=ARIAL,size=10); c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(r,1).alignment=Alignment(vertical='top'); ws.row_dimensions[r].height=max(15,13*(len(b)//110+1)); r+=1
ws.sheet_view.showGridLines=False
del wb['Sheet']
wb.save('Questoes_Mais_Erradas_Simulados_2026.xlsx')
# resumo p/ chat
print('CRÍTICAS <50%:',len(crit))
print('por área:',dict(byarea.most_common()))
print('por tipo:',dict(bycls.most_common()))
